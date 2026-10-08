from pathlib import Path
import hashlib, json, platform, importlib.metadata, zipfile
import fitz
from docx import Document
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[6]
OUT = Path(__file__).resolve().parent
def identity(p):
    return {'path': p.relative_to(ROOT).as_posix(), 'size_bytes': p.stat().st_size,
            'sha256': hashlib.sha256(p.read_bytes()).hexdigest()}
def main():
    raw = ROOT / 'runs/R19/inputs/raw'
    files = sorted(p for p in raw.iterdir() if p.is_file())
    facts = {'source_identity': [identity(p) for p in files], 'pdf': [], 'docx': {}, 'xlsx': {}, 'locks': []}
    for p in files:
        if p.suffix == '.pdf':
            with fitz.open(p) as d:
                for n, page in enumerate(d, 1):
                    facts['pdf'].append({'page': n, 'text': page.get_text(), 'images': len(page.get_images())})
                    page.get_pixmap(matrix=fitz.Matrix(1.6, 1.6), alpha=False).save(OUT / f'official-pdf-page-{n}.png')
        elif p.suffix == '.docx':
            d = Document(p)
            facts['docx']['paragraphs'] = [{'paragraph': n, 'text': x.text} for n, x in enumerate(d.paragraphs, 1)]
            facts['docx']['tables'] = [{'table': n, 'rows': [[c.text for c in r.cells] for r in t.rows]} for n, t in enumerate(d.tables, 1)]
            with zipfile.ZipFile(p) as z:
                facts['docx']['parts'] = z.namelist()
                facts['docx']['document_xml'] = z.read('word/document.xml').decode('utf-8')
        elif p.suffix == '.xlsx':
            w = load_workbook(p, data_only=False)
            wc = load_workbook(p, data_only=True)
            for s in w:
                facts['xlsx'][s.title] = {'dimension': s.calculate_dimension(), 'merged_ranges': [str(r) for r in s.merged_cells.ranges],
                    'cells': [{'cell': c.coordinate, 'value': c.value, 'cached': wc[s.title][c.coordinate].value,
                               'data_type': c.data_type, 'number_format': c.number_format}
                              for row in s for c in row if c.value is not None]}
            w.close(); wc.close()
    for loc in ['runs/R19/inputs/raw-lock.json', 'runs/R19/evaluation-lock.json', 'runs/R19/levels/L4/design-lock.json']:
        p = ROOT / loc
        lock = json.loads(p.read_text(encoding='utf-8-sig'))
        entries = []
        for expected in lock['files']:
            observed = identity(ROOT / expected['path'])
            entries.append({'expected': expected, 'observed': observed,
                            'match': all(observed[k] == v for k, v in expected.items())})
        facts['locks'].append({'lock': identity(p), 'entries': entries, 'all_match': all(e['match'] for e in entries)})
    facts['environment'] = {'python': platform.python_version(), 'libraries': {n: importlib.metadata.version(n) for n in ['PyMuPDF', 'python-docx', 'openpyxl', 'numpy', 'scipy', 'pypdf', 'pdfplumber']}}
    (OUT / 'source-audit.json').write_text(json.dumps(facts, ensure_ascii=False, indent=2), encoding='utf-8')
    text = '\n\n'.join(f"PDF PAGE {p['page']}\n{p['text']}" for p in facts['pdf'])
    text += '\n\n' + '\n'.join(f"DOCX P{x['paragraph']}: {x['text']}" for x in facts['docx']['paragraphs'])
    for t in facts['docx']['tables']:
        text += '\n' + '\n'.join(f"DOCX T{t['table']} R{n}: " + ' | '.join(row) for n, row in enumerate(t['rows'], 1))
    for name, s in facts['xlsx'].items():
        text += f"\nXLSX {name} {s['dimension']} merged={s['merged_ranges']}\n"
        text += '\n'.join(f"{c['cell']}: {c['value']!r} type={c['data_type']}" for c in s['cells'])
    (OUT / 'source-readback.txt').write_text(text, encoding='utf-8')
    print(text)
    print('LOCKS', [(x['lock']['path'], x['all_match']) for x in facts['locks']])
    print('ENVIRONMENT', facts['environment'])
if __name__ == '__main__':
    main()
