from pathlib import Path
import json, re
from docx import Document
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[6]
OUT=Path(__file__).resolve().parent
def main():
    d=Document(ROOT/'runs/R19/inputs/raw/附件1.docx')
    vehicles=[]
    for tid, lp, mp, cp in [('1',4,5,6),('2',8,9,10)]:
        dims=[float(v) for v in re.findall(r'(\d+)cm',d.paragraphs[lp-1].text)]
        mass=float(re.search(r'(\d+)kg',d.paragraphs[mp-1].text)[1])
        cost=float(re.search(r'(\d+)\s*元/趟',d.paragraphs[cp-1].text)[1])
        vehicles.append({'type_id':tid,'dims_cm':dims,'capacity_kg':mass,'cost_per_trip_yuan':cost,
                         'source':'附件1.docx','locators':{'dims':f'P{lp}','capacity':f'P{mp}','cost':f'P{cp}'}})
    cargo=[]
    for n,row in enumerate(d.tables[0].rows[1:],2):
        t,cl,ds,m,q=[c.text for c in row.cells]
        cargo.append({'type_id':t,'class':cl,'dims_cm':[float(v) for v in ds.split('×')],
                      'mass_kg':float(m),'quantity':int(q),'source':'附件1.docx','locator':f'T1 R{n}'})
    w=load_workbook(ROOT/'runs/R19/inputs/raw/附件2：验证数据集.xlsx',data_only=False)
    products=[]
    for r in range(3,11):
        entry={'product':w['箱装产品尺寸'][f'A{r}'].value,'row':r,'measurements':{}}
        for col in 'BCDEFG':
            c=w['箱装产品尺寸'][f'{col}{r}']; value=c.value
            nums=[] if value is None else [float(v) for v in re.findall(r'\d+(?:\.\d+)?',str(value))]
            entry['measurements'][c.coordinate]={'raw':value,'interval_mm':None if not nums else [min(nums),max(nums)],
                                                'interval_cm':None if not nums else [min(nums)/10,max(nums)/10],
                                                'origin':'laboratory' if col in 'BDF' else 'warehouse'}
        entry['note']=w['箱装产品尺寸'][f'H{r}'].value
        products.append(entry)
    w.close()
    contract={'purpose':'independently parsed raw fields, not answers or rankings','units':{'attachment1':'cm, kg, yuan/trip','attachment2_products':'mm','attachment2_vehicles':'m','area_conversion_cm2_to_m2':1/10000},
              'vehicles':vehicles,'cargo':cargo,'attachment2_products':products,
              'attachment2_missing':['product batch quantity','single mass','class/orientation constraints','vehicle weight capacity','distance and complete cost semantics'],
              'attachment2_vehicle_limits':{'E3':'公路单位运输费用元/1000km','F4:F6':'numeric values without independent header; do not treat as yuan/trip',
                                            'rows_12_19':'railing/nonuniform-floor geometry requires extra structure model; not closed rectangular inner height'}}
    (OUT/'source-contract.json').write_text(json.dumps(contract,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({'vehicles':vehicles,'cargo':cargo,'products':len(products)},ensure_ascii=False))
if __name__=='__main__':main()
