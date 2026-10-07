"""Source-to-output audit independent of normalization and pattern search."""
from pathlib import Path
import json,csv,math,re,subprocess,hashlib
from docx import Document
from validate import validate
R=Path(__file__).resolve().parents[1];A=R.parents[4];raw=A/'runs/R19/inputs/raw'
def rd(p):
 with Path(p).open(encoding='utf-8-sig') as f:return list(csv.DictReader(f))
doc=Document(raw/'附件1.docx');items=rd(R/'data/items.csv');v=json.loads((R/'data/vehicles.json').read_text(encoding='utf8'));source=[]
for row in doc.tables[0].rows[1:]:
 c=[x.text for x in row.cells];s={'type':c[0],'category':c[1],'dimensions':[int(x)*10 for x in c[2].split('×')],'weight':float(c[3]),'quantity':int(c[4])};source.append(s);rr=[r for r in items if r['cargo_type']==s['type']];assert len(rr)==s['quantity'];assert all(r['category']==s['category'] and float(r['weight_kg'])==s['weight'] and [float(r['canonical_'+k+'_mm']) for k in ['l','w','h']]==s['dimensions'] for r in rr)
assert len({r['item_id'] for r in items})==len(items)==3000
for j,offset in enumerate([3,7]):
 assert v[j]['dims']==[int(x)*10 for x in re.findall(r'(\d+)cm',doc.paragraphs[offset].text)]
 assert v[j]['payload_kg']==int(re.search(r'(\d+)kg',doc.paragraphs[offset+1].text).group(1));assert v[j]['cost_yuan_per_trip']==int(re.search(r'(\d+)\s*元',doc.paragraphs[offset+2].text).group(1))
audits=[]
for path in list((R/'results/final/selected').glob('*/run.json'))+list((R/'results/final/single').glob('*/run.json')):
 run=json.loads(path.read_text(encoding='utf8'));vv=json.loads((path.parent/'vehicles.json').read_text(encoding='utf8'));p=rd(path.parent/'placements.csv');z=validate(p,items,vv,full=run['full_inventory']);assert z['valid'],z['errors'];assert len(z['vehicles'])==run['vehicle_count'];assert sum(x['cost_yuan'] for x in z['vehicles'])==run['cost_yuan'];V={a['type_id']:a for a in vv};u=sum(x['volume_mm3'] for x in z['vehicles'])/sum(math.prod(V[x['vehicle_type']]['dims']) for x in z['vehicles']);w=sum(x['weight_kg'] for x in z['vehicles'])/sum(V[x['vehicle_type']]['payload_kg'] for x in z['vehicles']);assert abs(u-run['volume_utilization'])<1e-12 and abs(w-run['weight_utilization'])<1e-12;audits.append({'scenario':run['scenario'],'item_count':len(p),'all_checks_passed':True,'metrics_recomputed':True})
totalV=sum(s['quantity']*math.prod(s['dimensions']) for s in source);totalM=sum(s['quantity']*s['weight'] for s in source);lb=[max(math.ceil(totalV/(a['dims'][0]*a['dims'][1]*(a['dims'][2]-30))),math.ceil(totalM/a['payload_kg'])) for a in v];mix=[]
for a in range(50):
 for b in range(50):
  if sum(n*t['dims'][0]*t['dims'][1]*(t['dims'][2]-30) for n,t in zip([a,b],v))>=totalV and sum(n*t['payload_kg'] for n,t in zip([a,b],v))>=totalM:mix.append((a+b,a*450+b*700))
out={'source_rows':source,'source_reconciliation':True,'total_volume_mm3':totalV,'total_weight_kg':totalM,'independently_recomputed_bounds':lb+[min(x[0] for x in mix),min(x[1] for x in mix)],'runs':audits,'checker_independence':'separate code, same author context','model':None,'tokens':None,'cost':None};(R/'checks/source_result_audit.json').write_text(json.dumps(out,ensure_ascii=False,indent=2),encoding='utf8')
# Actual CLI rejection: mutate a genuine smoke coordinate beyond the roof.
p=rd(R/'checks/smoke/placements.csv');p[0]['z_mm']='10000';f=R/'checks/rejected_placements.csv'
with f.open('w',encoding='utf-8-sig',newline='') as stream:w=csv.DictWriter(stream,fieldnames=list(p[0]));w.writeheader();w.writerows(p)
cmd=['py','-3.12','-X','utf8','-B',str(R/'code/validate.py'),'--placements',str(f),'--items',str(R/'data/items.csv'),'--vehicles',str(R/'data/vehicles.json'),'--partial','--output',str(R/'checks/cli_rejection.json')];c=subprocess.run(cmd,capture_output=True,text=True,encoding='utf8');assert c.returncode==2;(R/'logs/cli_rejection_receipt.json').write_text(json.dumps({'command':cmd,'exit_code':c.returncode,'stdout':c.stdout,'stderr':c.stderr},ensure_ascii=False,indent=2),encoding='utf8')
print(json.dumps({'source_reconciliation':True,'recomputed_runs':len(audits),'bounds':out['independently_recomputed_bounds'],'actual_cli_rejection':c.returncode},ensure_ascii=False))
