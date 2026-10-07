from pathlib import Path
import json, hashlib, importlib.util, platform, sys, zipfile, xml.etree.ElementTree as ET
import pdfplumber
from openpyxl import load_workbook
base=Path.cwd(); out=base/'runs/R19/levels/L2/review/preparation'; raw=base/'runs/R19/inputs/raw'
record={'python':sys.version,'platform':platform.platform(),'model':None,'tokens':None,'cost':None,'raw':[]}
for p in sorted(raw.iterdir()):
 b=p.read_bytes(); record['raw'].append({'path':p.relative_to(base).as_posix(),'size_bytes':len(b),'sha256':hashlib.sha256(b).hexdigest()})
pdf=next(raw.glob('*.pdf'))
with pdfplumber.open(pdf) as doc:
 record['pdf']=[{'page':i+1,'width':p.width,'height':p.height,'text':p.extract_text(layout=False)} for i,p in enumerate(doc.pages)]
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
with zipfile.ZipFile(raw/'附件1.docx') as z:
 x=ET.fromstring(z.read('word/document.xml')); record['docx']={'paragraphs':[],'tables':[]}
 for i,p in enumerate(x.findall('.//w:body/w:p',ns)):
  record['docx']['paragraphs'].append({'index':i+1,'text':''.join(p.itertext()) if False else ''.join(t.text or '' for t in p.findall('.//w:t',ns))})
 for i,t in enumerate(x.findall('.//w:tbl',ns)):
  record['docx']['tables'].append({'index':i+1,'rows':[[''.join(n.text or '' for n in c.findall('.//w:t',ns)) for c in r.findall('w:tc',ns)] for r in t.findall('w:tr',ns)]})
xlsx=next(raw.glob('*.xlsx')); wb=load_workbook(xlsx,data_only=False); wd=load_workbook(xlsx,data_only=True)
record['xlsx']={'sheets':[],'named_ranges':str(list(wb.defined_names)),'external_links':len(wb._external_links)}
for s in wb:
 record['xlsx']['sheets'].append({'title':s.title,'max_row':s.max_row,'max_column':s.max_column,'merged':list(map(str,s.merged_cells.ranges)),'hidden':s.sheet_state,'cells':[{'cell':c.coordinate,'value':c.value,'data_type':c.data_type,'cached_value':wd[s.title][c.coordinate].value,'number_format':c.number_format} for row in s for c in row if c.value is not None]})
record['libraries']={k:bool(importlib.util.find_spec(k)) for k in ['fitz','pypdf','pdfplumber','openpyxl','docx','numpy','scipy','pandas','matplotlib']}
(out/'official-source-extract.json').write_text(json.dumps(record,ensure_ascii=False,indent=2,default=str),encoding='utf-8')
(out/'official-pdf-text.txt').write_text('\n\n'.join('PAGE '+str(p['page'])+'\n'+p['text'] for p in record['pdf']),encoding='utf-8')
print(json.dumps({'pdf':record['pdf'],'docx':record['docx'],'libraries':record['libraries'],'xlsx_sheets':[{k:v for k,v in s.items() if k!='cells'} for s in record['xlsx']['sheets']]},ensure_ascii=False,indent=2))
if record['libraries']['fitz']:
 import fitz
 with fitz.open(pdf) as doc:
  for i,p in enumerate(doc):p.get_pixmap(matrix=fitz.Matrix(1.6,1.6)).save(out/f'official-pdf-page-{i+1}.png')
 print('Rendered all PDF pages using existing PyMuPDF.')
