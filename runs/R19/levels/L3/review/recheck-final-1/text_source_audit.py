"""Own raw-read, Markdown/PDF semantic coverage and table-change audit."""
from pathlib import Path
import json,re,hashlib,math,time
import fitz
from docx import Document
import openpyxl
from PIL import Image,ImageDraw
H=Path(__file__).resolve().parent;B=H.parents[5];L=H.parents[1];N=L/'execution-v2';O=L/'execution'
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
def save(p,x):p.write_text(json.dumps(x,ensure_ascii=False,indent=2,allow_nan=False)+'\n',encoding='utf-8')
t=time.perf_counter();lock_rows=[]
for rel in ['runs/R19/candidate/C11-lock.json','runs/R19/inputs/raw-lock.json','runs/R19/levels/L3/design-lock.json','runs/R19/evaluation-lock.json']:
    j=json.loads((B/rel).read_text(encoding='utf-8-sig'))
    for row in j['files']:
        p=B/row['path'];lock_rows.append({'lock':rel,'path':row['path'],'matches':sha(p)==row['sha256']})
assert all(x['matches'] for x in lock_rows)
raw=B/'runs/R19/inputs/raw'
d=Document(raw/'附件1.docx')
docx={'paragraphs':[{'paragraph':i+1,'text':p.text} for i,p in enumerate(d.paragraphs)],'tables':[[[c.text for c in row.cells] for row in table.rows] for table in d.tables]}
wb=openpyxl.load_workbook(raw/'附件2：验证数据集.xlsx',data_only=False,read_only=False)
xlsx={s.title:[{'cell':c.coordinate,'value':c.value,'type':c.data_type} for row in s for c in row if c.value is not None] for s in wb}
pdf=fitz.open(raw/'2026年第十六届MathorCup数学应用挑战赛题目—D题.pdf')
save(H/'raw-reread.json',{'source_locks':lock_rows,'source_count':len(lock_rows),'official_pdf':[p.get_text() for p in pdf],'docx':docx,'xlsx':xlsx})
def tables(text):
    groups=[];current=[]
    for line in text.splitlines()+['']:
        if line.strip().startswith('|'):
            cells=[s.strip() for s in line.strip().strip('|').split('|')]
            if not all(re.fullmatch('[-: ]+',s) for s in cells):current.append(cells)
        elif current:groups.append(current);current=[]
    return groups
def units(text):
    result=[]
    for line in text.splitlines():
        s=line.strip()
        if not s or s.startswith('```'):continue
        if s.startswith('|'):
            cells=[c.strip() for c in s.strip('|').split('|')]
            if not all(re.fullmatch('[-: ]+',c) for c in cells):result.extend(cells)
        elif s.startswith('!['):result.append(re.fullmatch(r'!\[(.*?)\]\((.*?)\)',s)[1])
        else:result.append(re.sub(r'^#{1,3} ','',s))
    return result
norm=lambda s:re.sub(r'\s+','',s.replace('**',''))
info={};science_crops=[];visual=H/'pdf_visual'
for name in ['paper','technical_report']:
    md=(N/f'{name}.md').read_text(encoding='utf-8-sig')
    old=(O/f'{name}.md').read_text(encoding='utf-8-sig')
    doc=fitz.open(N/f'{name}.pdf');body=''
    for page in doc:
        # Strip only reviewer-observed fixed footer region, not manuscript text.
        body+=page.get_text(clip=fitz.Rect(0,0,page.rect.width,page.rect.height-35))
    stream=norm(body);cursor=0;errors=[];display=units(md)
    for i,s in enumerate(display):
        found=stream.find(norm(s),cursor)
        if found<0:errors.append({'index':i,'source':s})
        else:cursor=found+len(norm(s))
    newtab=tables(md);oldtab=tables(old)
    # Rows tied to original independently checked formal and scenario records.
    key=lambda c:c[0] in ['Q1-S1','Q1-S2','Q1-F1','Q1-F2','Q2-N','Q2-C'] or c[0]=='base' or '_' in c[0]
    current=[r for table in newtab for r in table if key(r)]
    previous=[r for table in oldtab for r in table if key(r)]
    changed=[]
    for row in current:
        if row not in previous:changed.append(row)
    for pi,page in enumerate(doc):
        for bi,block in enumerate(page.get_text('dict')['blocks']):
            for li,line in enumerate(block.get('lines',[])):
                s=''.join(span['text'] for span in line['spans'])
                if any(ch in s for ch in '²³⁻'):
                    box=fitz.Rect(line['bbox']);crop=fitz.Rect(40,max(0,box.y0-8),page.rect.width-40,min(page.rect.height,box.y1+8))
                    p=visual/f'science-{name}-{pi+1:02}-{bi}-{li}.png';page.get_pixmap(matrix=fitz.Matrix(1.7,1.7),clip=crop).save(p)
                    science_crops.append({'document':name,'page':pi+1,'text':s,'bbox':list(crop),'path':str(p)})
    info[name]={'display_units':len(display),'ordered_pdf_semantic_errors':errors,'table_rows_compared_to_initial':len(current),'changed_table_rows':changed,'science_char_counts':{c:{'md':md.count(c),'pdf':body.count(c)} for c in '²³⁻'},'pages':len(doc)}
assert all(not x['ordered_pdf_semantic_errors'] and not x['changed_table_rows'] for x in info.values()),info
for start in range(0,len(science_crops),7):
    pieces=[]
    for row in science_crops[start:start+7]:
        im=Image.open(row['path']).convert('RGB');canvas=Image.new('RGB',(im.width,max(85,im.height+26)),'#eeeeee');canvas.paste(im,(0,22));ImageDraw.Draw(canvas).text((5,3),f"{row['document']} p{row['page']}",fill='black');pieces.append(canvas)
    sheet=Image.new('RGB',(max(x.width for x in pieces),sum(x.height for x in pieces)+6*len(pieces)),'white');y=0
    for im in pieces:sheet.paste(im,(0,y));y+=im.height+6
    sheet.save(visual/f'science-contact-{start//7+1}.png')
cargo=d.tables[0].rows[1:]
total_volume=sum(math.prod(float(v) for v in row.cells[2].text.split('×'))*int(row.cells[4].text) for row in cargo)/1e6
total_weight=sum(float(row.cells[3].text)*int(row.cells[4].text) for row in cargo)
derived={'total_items':sum(int(row.cells[4].text) for row in cargo),'volume_m3':total_volume,'weight_kg':total_weight,'effective_volume_m3':[420*210*217/1e6,680*245*247/1e6],'density_capacity_bounds_percent':[187.5*420*210*217/1e6/6000*100,187.5*680*245*247/1e6/10000*100],'gap_relative_to_lower_percent':[(20-16)/16*100,(10-7)/7*100,(7000-4900)/4900*100]}
save(H/'text-and-source-audit.json',{'documents':info,'all_display_units':sum(x['display_units'] for x in info.values()),'science_crop_lines':science_crops,'derived_from_raw':derived,'elapsed_seconds':time.perf_counter()-t,'limits':['Ordered textual coverage supports faithful transcription, not automatic proof of readable Chinese argument or scientific correctness.','All actual page/crop visual observations are recorded separately.']})
print(json.dumps({'source_locks':len(lock_rows),'document_checks':info,'crop_lines':len(science_crops),'derived':derived,'elapsed_seconds':time.perf_counter()-t},ensure_ascii=False))
