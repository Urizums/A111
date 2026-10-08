"""Independent original-source audit; writes only beside this script; no solving."""
import hashlib, importlib, json, os, platform, sys, time
from pathlib import Path
from datetime import datetime, timezone

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[5]
def save(name, value):
    (HERE/name).write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding='utf-8')

start = time.perf_counter()
probe = {'utc': datetime.now(timezone.utc).isoformat(), 'python': sys.version,
         'executable': sys.executable, 'platform': platform.platform(),
         'logical_cpus': os.cpu_count(), 'model': None, 'token_usage': None, 'cost': None,
         'modules': {}}
for name in ['numpy','scipy','pandas','matplotlib','openpyxl','docx','pypdf','fitz','psutil']:
    try:
        module = importlib.import_module(name)
        probe['modules'][name] = {'available': True, 'version': getattr(module, '__version__', None)}
    except Exception as exc:
        probe['modules'][name] = {'available': False, 'error': repr(exc)}
if probe['modules']['psutil']['available']:
    import psutil
    probe['memory_available_bytes'] = psutil.virtual_memory().available
    probe['physical_cpus'] = psutil.cpu_count(logical=False)
    probe['affinity'] = psutil.Process().cpu_affinity()
save('resource-probe.json', probe)

locks = ['runs/R19/candidate/C11-lock.json', 'runs/R19/inputs/raw-lock.json',
         'runs/R19/levels/L3/design-lock.json', 'runs/R19/evaluation-lock.json']
verification = []
for lock in locks:
    data = json.loads((ROOT/lock).read_text(encoding='utf-8'))
    for item in data['files']:
        path = ROOT/item['path']
        blob = path.read_bytes()
        actual = hashlib.sha256(blob).hexdigest()
        verification.append({'lock': lock, 'path': item['path'], 'bytes': len(blob),
                             'expected_sha256': item['sha256'], 'actual_sha256': actual,
                             'matches': actual == item['sha256'] and len(blob) == item.get('size_bytes', len(blob))})
save('input-lock-verification.json', verification)
assert all(item['matches'] for item in verification), 'Source lock mismatch'

import fitz
from docx import Document
from openpyxl import load_workbook
raw = ROOT/'runs/R19/inputs/raw'
pdf = next(raw.glob('*.pdf'))
pages = []
with fitz.open(pdf) as document:
    for i, page in enumerate(document):
        pages.append({'page': i+1, 'text': page.get_text()})
        page.get_pixmap(matrix=fitz.Matrix(1.5,1.5)).save(HERE/f'raw-pdf-page-{i+1}.png')
doc = Document(raw/'附件1.docx')
paragraphs = [{'paragraph': i+1, 'text': p.text} for i,p in enumerate(doc.paragraphs)]
tables = [{'table': i+1, 'rows': [[c.text for c in r.cells] for r in t.rows]} for i,t in enumerate(doc.tables)]
wb = load_workbook(raw/'附件2：验证数据集.xlsx', data_only=False)
sheets = []
for ws in wb:
    sheets.append({'sheet': ws.title, 'max_row': ws.max_row, 'max_column': ws.max_column,
                   'merged_ranges': [str(r) for r in ws.merged_cells.ranges],
                   'cells': [{'address': c.coordinate,'value':c.value,'data_type':c.data_type}
                             for row in ws for c in row if c.value is not None]})
audit = {'role': '/root/r19_l3_reviewer', 'phase': 'preparation_only', 'pdf': pages,
         'docx': {'paragraphs': paragraphs,'tables':tables}, 'xlsx': sheets,
         'elapsed_seconds': time.perf_counter()-start, 'production_inputs_read':False}
save('independent-source-audit.json', audit)
print(json.dumps({'hash_entries':len(verification),'all_match':True,'pdf_pages':len(pages),
                  'docx_paragraphs':len(paragraphs),'xlsx_sheets':len(sheets),
                  'elapsed_seconds':audit['elapsed_seconds']},ensure_ascii=False))
