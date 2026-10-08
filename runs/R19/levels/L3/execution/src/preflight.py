from pathlib import Path
import csv,json,hashlib,sys,importlib,re,datetime
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader
E=Path(__file__).resolve().parents[1]
ROOT=E.parents[2]
def save(p,x):
 p.parent.mkdir(parents=True,exist_ok=True);p.write_text(json.dumps(x,ensure_ascii=False,indent=2),encoding='utf-8')
def main():
 raw=ROOT/'inputs/raw';lock=json.loads((ROOT/'inputs/raw-lock.json').read_text(encoding='utf-8'))
 audit=[]
 for r in lock['files']:
  p=raw/Path(r['path']).name;h=hashlib.sha256(p.read_bytes()).hexdigest();assert h==r['sha256'];audit.append({'file':p.name,'sha256':h,'bytes':p.stat().st_size,'match':True})
 doc=Document(raw/'附件1.docx');table=[[c.text for c in row.cells] for row in doc.tables[0].rows]
 paras=[{'paragraph':i+1,'text':p.text} for i,p in enumerate(doc.paragraphs)]
 sheets={}
 wb=load_workbook(raw/'附件2：验证数据集.xlsx',data_only=False)
 for s in wb:
  sheets[s.title]={'merged':[str(x) for x in s.merged_cells.ranges],'cells':[{'cell':c.coordinate,'value':c.value,'type':c.data_type} for row in s for c in row if c.value is not None]}
 pdf=[{'page':i+1,'text':p.extract_text()} for i,p in enumerate(PdfReader(raw/'2026年第十六届MathorCup数学应用挑战赛题目—D题.pdf').pages)]
 save(E/'data/raw_audit.json',{'files':audit,'docx_paragraphs':paras,'docx_table':table,'xlsx':sheets,'pdf':pdf})
 vals=[('G1','standard',60,40,30,12,800),('G2','standard',50,35,25,8,1000),('G3','fragile',70,50,40,15,300),('G4','directional',80,60,50,25,400),('G5','directional',40,40,60,18,500)]
 cargo=[]
 for j,t in enumerate(vals):
  text=' '.join(table[j+1]);nums=[int(n) for n in re.findall(r'\d+',text)]
  # Official type ID contributes its first digit; dimensions/mass/quantity follow.
  assert nums[-5:]==list(t[2:]),(text,nums,t)
  cargo.append(dict(zip(['cargo_type','category','l_cm','w_cm','h_cm','weight_kg','quantity'],t),source=f'附件1 表1 R{j+2}',origin='official'))
 vehicles=[{'vehicle_type':'T1','l_cm':420,'w_cm':210,'h_cm':220,'capacity_kg':6000,'cost_yuan':450,'source':'附件1 P4-P6','origin':'official'}, {'vehicle_type':'T2','l_cm':680,'w_cm':245,'h_cm':250,'capacity_kg':10000,'cost_yuan':700,'source':'附件1 P8-P10','origin':'official'}]
 for p in vehicles:
  numbers=[p[k] for k in ['l_cm','w_cm','h_cm','capacity_kg','cost_yuan']]
  joined=' '.join(x['text'] for x in paras[3:6] if p['vehicle_type']=='T1') if p['vehicle_type']=='T1' else ' '.join(x['text'] for x in paras[7:10])
  assert all(str(n) in joined for n in numbers)
 for name,rows in [('cargo',cargo),('vehicles',vehicles)]:
  with (E/f'data/{name}.csv').open('w',newline='',encoding='utf-8-sig') as f:
   wr=csv.DictWriter(f,fieldnames=list(rows[0]));wr.writeheader();wr.writerows(rows)
 save(E/'data/instance.json',{'cargo':cargo,'vehicles':vehicles})
 save(E/'data/assumptions.json',{'coordinate':'right-rear-bottom; x forward, y left, z up','units':{'length':'cm','mass':'kg','cost':'yuan/trip'},'tolerance_cm':1e-7,'top_gap_cm':3,'pressure_limit_kg_m2':500,'orientation':'standard and fragile: all distinct orthogonal permutations; directional: original only','support':'non-floor complete bottom coverage by one lower box; fragile support must be standard; nothing above fragile','load':'cumulative upper mass/contact area, own mass excluded at box top; single-parent vertical chains preserve mass','gravity':'uniform mass, geometric center; contained footprints guarantee directional center condition','scenario_variants':['fragile upright','fragile floor only','pressure own mass included','pressure 300/700','multi-support not exploited'],'independence':'producer self-check; external review pending','unknown_model':None,'unknown_tokens':None,'unknown_cost':None})
 env={'python':sys.version,'libraries':{m:getattr(importlib.import_module(m),'__version__','available') for m in ['numpy','scipy','pandas','matplotlib','docx','openpyxl','pypdf','fitz','reportlab']}}
 save(E/'logs/environment.json',env)
 save(E/'logs/resources.json',{'declared_at':datetime.datetime.now(datetime.UTC).isoformat(),'phase':'compute-preflight','window_seconds':3600,'process_memory_target_MB':1800,'parallel_solver_processes':1,'enforcement':'internal phase/run deadlines; memory observations recorded, target not OS enforced','model':None,'tokens':None,'cost':None,'network':False,'installation':False,'submission':False})
 save(E/'data/contracts.json',{'tasks':{'Q1-S1':'T1 finite inventory subset; nondominated volume/mass','Q1-S2':'T2 finite inventory subset; nondominated volume/mass','Q1-F1':'all items; T1 only; minimize number','Q1-F2':'all items; T2 only; minimize number','Q2-N':'all items; T1/T2; minimize count then cost','Q2-C':'all items; T1/T2; minimize cost then count'},'total_volume_m3':sum(c['l_cm']*c['w_cm']*c['h_cm']*c['quantity'] for c in cargo)/1e6,'total_weight_kg':sum(c['weight_kg']*c['quantity'] for c in cargo),'total_items':sum(c['quantity'] for c in cargo)})
 print(json.dumps({'source_checks':audit,'cargo':cargo,'vehicles':vehicles,'contracts':json.loads((E/'data/contracts.json').read_text(encoding='utf-8'))},ensure_ascii=False))
if __name__=='__main__':main()
