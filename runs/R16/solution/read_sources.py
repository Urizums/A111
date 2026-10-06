from pathlib import Path
from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook

R16 = Path(__file__).resolve().parents[1]
SRC = R16 / 'source' / 'official_extracted' / 'D_corrected'
OUT = R16 / 'solution' / 'source_readout.txt'
lines = []
pdf = PdfReader(str(next(SRC.glob('*.pdf'))))
lines.append(f'PDF pages={len(pdf.pages)}')
for i, page in enumerate(pdf.pages):
    lines.append(f'--- PDF PAGE {i+1} ---\n{page.extract_text()}')
doc = Document(SRC / '附件1.docx')
lines.append('--- ATTACHMENT 1 DOCX PARAGRAPHS ---')
lines.extend(p.text for p in doc.paragraphs)
for i, table in enumerate(doc.tables):
    lines.append(f'--- DOCX TABLE {i+1} ---')
    lines.extend(' | '.join(c.text for c in row.cells) for row in table.rows)
wb = load_workbook(SRC / '附件2：验证数据集.xlsx', data_only=True)
lines.append(f'--- ATTACHMENT 2 XLSX sheets={wb.sheetnames} ---')
for ws in wb.worksheets:
    lines.append(f'--- SHEET {ws.title} {ws.max_row}x{ws.max_column} ---')
    for row in ws.iter_rows(values_only=True):
        lines.append(' | '.join('' if v is None else str(v) for v in row))
OUT.write_text('\n'.join(lines), encoding='utf-8')
print(OUT)
print('\n'.join(lines))
